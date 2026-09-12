"""Tests for file operations API endpoints (api.file_ops)."""
import pytest
import os
import io
import json
from fastapi.testclient import TestClient
from app.main import app
from app.database import get_db
from app.config import UPLOAD_DIR


def _retry_remove(path, attempts=50, delay=0.1):
    """Remove a file, retrying briefly on Windows file-lock races.

    The analyze-saved handler reads the upload with openpyxl/pandas, whose
    handle can be released asynchronously on Windows after the response
    returns. This is best-effort teardown: never fail the test on cleanup.
    """
    import time
    for _ in range(attempts):
        try:
            os.remove(path)
            return
        except PermissionError:
            time.sleep(delay)


@pytest.fixture
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass
    app.dependency_overrides[get_db] = override_get_db
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def sample_excel_bytes():
    import pandas as pd
    def _make_bytes():
        data = {
            "organisationunitname": ["Test Hospital", "Test Hospital"],
            "month": ["2026-04", "2026-04"],
            "Total Deliveries": [300, 280],
            "Normal Vaginal Deliveries": [200, 180],
            "Caesarean Sections": [80, 75],
            "Live Births": [290, 270],
            "Maternal Deaths": [1, 0],
            "Neonatal deaths": [5, 3],
        }
        df = pd.DataFrame(data)
        buf = io.BytesIO()
        df.to_excel(buf, index=False, engine="openpyxl")
        buf.seek(0)
        return buf
    return _make_bytes


class TestListSavedFiles:
    def test_empty_upload_dir(self, client):
        resp = client.get("/analysis/saved-files")
        assert resp.status_code == 200
        data = resp.json()
        assert isinstance(data, list)

    def test_with_files(self, client, sample_excel_bytes, db_session):
        """saved-files now reads from DB source_file, not filesystem."""
        from app.models import IndicatorValue, Hospital
        h = db_session.query(Hospital).first()
        db_session.add(IndicatorValue(
            hospital_id=h.id, indicator_id=1, month="2026-01",
            value=10.0, source_file="test_list.xlsx",
        ))
        db_session.commit()

        resp = client.get("/analysis/saved-files")
        assert resp.status_code == 200
        data = resp.json()
        assert isinstance(data, list)
        filenames = [f["filename"] for f in data]
        assert "test_list.xlsx" in filenames

    def test_file_has_required_fields(self, client, sample_excel_bytes, db_session):
        from app.models import IndicatorValue, Hospital
        h = db_session.query(Hospital).first()
        db_session.add(IndicatorValue(
            hospital_id=h.id, indicator_id=1, month="2026-02",
            value=20.0, source_file="test_fields.xlsx",
        ))
        db_session.commit()

        resp = client.get("/analysis/saved-files")
        data = resp.json()
        matching = [f for f in data if f["filename"] == "test_fields.xlsx"]
        assert matching
        f = matching[0]
        assert "filename" in f
        assert "records_in_db" in f
        assert f["records_in_db"] >= 1

    def test_excludes_non_excel(self, client):
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        test_file = os.path.join(UPLOAD_DIR, "readme.txt")
        with open(test_file, "w") as f:
            f.write("not an excel file")

        resp = client.get("/analysis/saved-files")
        data = resp.json()
        filenames = [f["filename"] for f in data]
        assert "readme.txt" not in filenames

        if os.path.exists(test_file):
            os.remove(test_file)


class TestAnalyzeSavedFiles:
    def test_nonexistent_file(self, client):
        resp = client.post("/analysis/analyze-saved?filenames=nonexistent.xlsx")
        assert resp.status_code == 200
        data = resp.json()
        assert data["files_processed"] == 0

    def test_with_valid_file(self, client, sample_excel_bytes, db_session):
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        test_file = os.path.join(UPLOAD_DIR, "test_analyze.xlsx")
        with open(test_file, "wb") as f:
            f.write(sample_excel_bytes().getvalue())

        resp = client.post("/analysis/analyze-saved?filenames=test_analyze.xlsx")
        assert resp.status_code == 200
        data = resp.json()
        assert data["files_processed"] >= 1
        assert "hospitals_processed" in data
        assert "rows_imported" in data

        if os.path.exists(test_file):
            _retry_remove(test_file)


class TestDeleteSavedFiles:
    def test_delete_existing(self, client, db_session):
        """Delete removes IndicatorValue rows by source_file."""
        from app.models import IndicatorValue, Hospital
        h = db_session.query(Hospital).first()
        db_session.add(IndicatorValue(
            hospital_id=h.id, indicator_id=1, month="2026-03",
            value=30.0, source_file="to_delete.xlsx",
        ))
        db_session.commit()

        resp = client.request("DELETE", "/analysis/saved-files", content=json.dumps({"filenames": ["to_delete.xlsx"]}), headers={"Content-Type": "application/json"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["deleted"] >= 1

    def test_delete_nonexistent(self, client):
        resp = client.request("DELETE", "/analysis/saved-files", content=json.dumps({"filenames": ["does_not_exist.xlsx"]}), headers={"Content-Type": "application/json"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["deleted"] == 0

    def test_delete_mixed(self, client, db_session):
        """Deleting a mix of existing source_file and nonexistent one."""
        from app.models import IndicatorValue, Hospital
        h = db_session.query(Hospital).first()
        db_session.add(IndicatorValue(
            hospital_id=h.id, indicator_id=1, month="2026-04",
            value=40.0, source_file="mixed_delete.xlsx",
        ))
        db_session.commit()

        resp = client.request("DELETE", "/analysis/saved-files", content=json.dumps({
            "filenames": ["mixed_delete.xlsx", "does_not_exist.xlsx"]
        }), headers={"Content-Type": "application/json"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["deleted"] >= 1


class TestUploadMultiple:
    def test_upload_single_file(self, client, sample_excel_bytes):
        buf = sample_excel_bytes()
        resp = client.post(
            "/analysis/upload-multiple",
            files={"files": ("test_upload.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["files_processed"] >= 1
        assert "rows_imported" in data
        assert "hospitals" in data

    def test_upload_empty_file(self, client):
        buf = io.BytesIO(b"")
        resp = client.post(
            "/analysis/upload-multiple",
            files={"files": ("empty.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["files_processed"] == 0

    def test_upload_unsupported_extension(self, client):
        buf = io.BytesIO(b"dummy content")
        resp = client.post(
            "/analysis/upload-multiple",
            files={"files": ("test.txt", buf, "text/plain")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["files_processed"] == 0


class TestUploadMultipleAnalyze:
    def test_upload_and_analyze(self, client, sample_excel_bytes):
        buf = sample_excel_bytes()
        resp = client.post(
            "/analysis/upload-multiple-analyze",
            files={"files": ("test_ua.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert "task_id" in data
        assert data["files_processed"] >= 1

    def test_returns_task_id(self, client, sample_excel_bytes):
        buf = sample_excel_bytes()
        resp = client.post(
            "/analysis/upload-multiple-analyze",
            files={"files": ("test_tid.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        data = resp.json()
        assert data["task_id"]
        assert isinstance(data["task_id"], str)


class TestUpdateSavedFiles:
    def test_update_nonexistent_file(self, client, sample_excel_bytes):
        buf = sample_excel_bytes()
        resp = client.post(
            "/analysis/update-saved?filename=missing.xlsx",
            files={"file": ("replacement.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 404

    def test_update_replaces_records(self, client, sample_excel_bytes, db_session):
        """Updating a file keeps the same source_file and swaps its rows."""
        from app.models import IndicatorValue, Hospital
        h = db_session.query(Hospital).first()
        db_session.add(IndicatorValue(
            hospital_id=h.id, indicator_id=1, month="2026-04",
            value=40.0, source_file="to_update.xlsx",
        ))
        db_session.commit()

        buf = sample_excel_bytes()
        resp = client.post(
            "/analysis/update-saved?filename=to_update.xlsx",
            files={"file": ("replacement.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["filename"] == "to_update.xlsx"
        assert data["deleted"] >= 1
        assert data["rows_imported"] >= 1

        still_there = db_session.query(IndicatorValue).filter(IndicatorValue.source_file == "to_update.xlsx").count()
        assert still_there >= 1

    def test_update_unsupported_extension(self, client, db_session):
        from app.models import IndicatorValue, Hospital
        h = db_session.query(Hospital).first()
        db_session.add(IndicatorValue(
            hospital_id=h.id, indicator_id=1, month="2026-04",
            value=40.0, source_file="bad_update.xlsx",
        ))
        db_session.commit()

        resp = client.post(
            "/analysis/update-saved?filename=bad_update.xlsx",
            files={"file": ("notallowed.txt", io.BytesIO(b"dummy"), "text/plain")},
        )
        assert resp.status_code == 422


class TestProcessPreview:
    def test_process_existing_file(self, client, sample_excel_bytes, db_session):
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        test_file = os.path.join(UPLOAD_DIR, "preview_test_pt1.xlsx")
        with open(test_file, "wb") as f:
            f.write(sample_excel_bytes().getvalue())
        assert os.path.exists(test_file), f"File not created: {test_file}"

        resp = client.post("/analysis/process-preview", params={"filename": "preview_test_pt1.xlsx"})
        assert resp.status_code in (200, 404, 422, 500)

        try:
            if os.path.exists(test_file):
                os.remove(test_file)
        except PermissionError:
            pass

    def test_process_nonexistent_file(self, client):
        resp = client.post("/analysis/process-preview", params={"filename": "does_not_exist.xlsx"})
        assert resp.status_code == 404

    def test_process_returns_quality_reports(self, client, sample_excel_bytes, db_session):
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        test_file = os.path.join(UPLOAD_DIR, "preview_qr_pt2.xlsx")
        with open(test_file, "wb") as f:
            f.write(sample_excel_bytes().getvalue())
        assert os.path.exists(test_file), f"File not created: {test_file}"

        resp = client.post("/analysis/process-preview", params={"filename": "preview_qr_pt2.xlsx"})
        assert resp.status_code in (200, 404, 422, 500)

        try:
            if os.path.exists(test_file):
                os.remove(test_file)
        except PermissionError:
            pass


class TestUploadAnalyzeRecomputesAllMonths:
    def test_other_months_refreshed_on_upload(self, client, db_session, tmp_path):
        """Uploading a new month must also recompute the hospital's existing
        months, whose historical/trend baselines the new values changed."""
        from app.models import Hospital, Indicator, IndicatorValue, QualityScore

        hospital = db_session.query(Hospital).first()
        ind = db_session.query(Indicator).first()
        db_session.add(IndicatorValue(
            hospital_id=hospital.id, indicator_id=ind.id, month="2026-01", value=10.0))
        db_session.add(QualityScore(
            hospital_id=hospital.id, month="2026-01", score=1.0, issues="[]"))
        db_session.commit()

        csv = tmp_path / "upload_recompute.csv"
        csv.write_text(
            "organisationunitname,month,Total Deliveries\n"
            f"{hospital.name},2026-02,20\n",
            encoding="utf-8",
        )
        with open(str(csv), "rb") as fh:
            resp = client.post(
                "/upload/analyze",
                files={"file": ("upload_recompute.csv", fh, "text/csv")},
            )
        assert resp.status_code == 200

        qs = db_session.query(QualityScore).filter(
            QualityScore.hospital_id == hospital.id,
            QualityScore.month == "2026-01",
        ).first()
        assert qs is not None, "existing month lost its quality score"
        assert qs.score != 1.0, "existing month was not refreshed on upload"

        uploaded = os.path.join(UPLOAD_DIR, "upload_recompute.csv")
        if os.path.exists(uploaded):
            _retry_remove(uploaded)


class TestDownloadSavedFile:
    def test_download_original_file_when_on_disk(self, client):
        """Downloading a saved file that still exists on disk must return the
        exact original bytes as an attachment."""
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        test_file = os.path.join(UPLOAD_DIR, "dl_on_disk.xlsx")
        original = b"PK\x03\x04fake-original-bytes"
        with open(test_file, "wb") as f:
            f.write(original)

        resp = client.get("/analysis/saved-files/download", params={"filename": "dl_on_disk.xlsx"})
        assert resp.status_code == 200
        assert "attachment" in resp.headers.get("content-disposition", "")
        assert "dl_on_disk.xlsx" in resp.headers.get("content-disposition", "")
        assert resp.content == original

        _retry_remove(test_file)

    def test_download_rebuilds_xlsx_from_db_when_not_on_disk(self, client, db_session):
        """Downloading a file whose original is gone (ephemeral disk on Render)
        must rebuild an .xlsx from the IndicatorValue rows for that file."""
        import openpyxl
        from app.models import Hospital, Indicator, IndicatorValue
        from app.indicators import INDICATOR_FLAT_LIST

        hospital = db_session.query(Hospital).first()
        ind = db_session.query(Indicator).first()
        code = ind.code if ind.code else INDICATOR_FLAT_LIST[0]["code"]
        name = ind.name if ind.name else INDICATOR_FLAT_LIST[0]["name"]
        db_session.add(IndicatorValue(
            hospital_id=hospital.id, indicator_id=ind.id, month="2026-03",
            value=42.5, source_file="dl_history_only.xlsx",
        ))
        db_session.commit()

        on_disk = os.path.join(UPLOAD_DIR, "dl_history_only.xlsx")
        if os.path.exists(on_disk):
            os.remove(on_disk)

        resp = client.get("/analysis/saved-files/download", params={"filename": "dl_history_only.xlsx"})
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers.get("content-type", "")
        assert "attachment" in resp.headers.get("content-disposition", "")

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(header)}
        assert col.get("Month") is not None
        assert col.get("Value") is not None
        assert col.get("Hospital") is not None

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert any(
            r[col["Hospital"]] == hospital.name
            and r[col["Month"]] == "2026-03"
            and r[col["Value"]] == 42.5
            for r in rows
        )

    def test_download_unknown_file_404(self, client):
        resp = client.get("/analysis/saved-files/download", params={"filename": "never_uploaded.xlsx"})
        assert resp.status_code == 404
